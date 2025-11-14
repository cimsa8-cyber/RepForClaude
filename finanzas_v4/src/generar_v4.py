# -*- coding: utf-8 -*-
"""
🏗️ GENERADOR DE EXCEL v4.0 LIMPIO - Sistema de Finanzas
=========================================================

Crea un archivo Excel v4.0 desde cero con:
- Estructura correcta de columnas
- Headers formateados
- Formatos de fecha/moneda configurados
- Fórmulas de CxP/CxC implementadas
- Datos de ejemplo (opcional)

Este script garantiza que el archivo base esté correcto desde el inicio.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

from config import (
    ESTRUCTURA_TRANSACCIONES,
    ESTRUCTURA_CXP,
    ESTRUCTURA_CXC,
    HOJAS_CONFIG,
    get_formula_cxp,
    get_formula_cxc,
    TRANSACCIONES_EJEMPLO,
    VERSION
)

# ==============================================================================
# ESTILO Y FORMATO
# ==============================================================================

def aplicar_estilo_header(ws, fila, color_hex):
    """Aplica estilo a la fila de headers"""
    fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type='solid')
    font = Font(bold=True, color='FFFFFF', size=11)
    alignment = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for cell in ws[fila]:
        if cell.value:
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment
            cell.border = border

def ajustar_anchos_columnas(ws, estructura):
    """Ajusta el ancho de las columnas según el contenido"""
    anchos = {
        'Fecha': 12,
        'Tipo': 12,
        'Categoría': 15,
        'Subcategoría': 15,
        'Descripción': 35,
        'Entidad': 20,
        'Cuenta': 18,
        'Método Pago': 15,
        'Monto': 15,
        'Referencia': 12,
        'Notas': 25,
        'Estado': 12,
        'IVA': 12,
        'Fecha Vencimiento': 18,
        'Número Factura': 15,
        'Cliente': 20,
        'Días Pendientes': 15
    }

    for nombre_campo, config in estructura.items():
        col_letra = config.get('col_letra') or get_column_letter(config['col'])
        ancho = anchos.get(nombre_campo, 15)
        ws.column_dimensions[col_letra].width = ancho

# ==============================================================================
# CREACIÓN DE HOJA: TRANSACCIONES
# ==============================================================================

def crear_hoja_transacciones(wb, incluir_ejemplos=True):
    """Crea la hoja de TRANSACCIONES con estructura correcta"""
    print("📝 Creando hoja TRANSACCIONES...")

    # Crear o obtener hoja
    if 'TRANSACCIONES' in wb.sheetnames:
        ws = wb['TRANSACCIONES']
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('TRANSACCIONES')

    # Escribir headers (fila 1)
    for nombre_campo, config in ESTRUCTURA_TRANSACCIONES.items():
        col = config['col']
        ws.cell(1, col, nombre_campo)

    # Aplicar estilo a headers
    aplicar_estilo_header(ws, 1, HOJAS_CONFIG['TRANSACCIONES']['color_header'])

    # Ajustar anchos de columna
    ajustar_anchos_columnas(ws, ESTRUCTURA_TRANSACCIONES)

    # Congelar primera fila
    ws.freeze_panes = 'A2'

    # Insertar datos de ejemplo si se solicita
    if incluir_ejemplos:
        print("  📊 Insertando datos de ejemplo...")
        for ejemplo in TRANSACCIONES_EJEMPLO:
            fila = ws.max_row + 1

            for nombre_campo, config in ESTRUCTURA_TRANSACCIONES.items():
                col = config['col']
                valor = ejemplo.get(nombre_campo)

                # Convertir fecha string a datetime
                if nombre_campo in ['Fecha', 'Fecha Vencimiento'] and isinstance(valor, str) and valor:
                    try:
                        valor = datetime.strptime(valor, '%d/%m/%Y')
                    except ValueError:
                        pass

                # Escribir valor
                celda = ws.cell(fila, col, valor)

                # Aplicar formato
                formato = config.get('formato_openpyxl')
                if formato:
                    celda.number_format = formato

    print("  ✅ Hoja TRANSACCIONES creada")
    return ws

# ==============================================================================
# CREACIÓN DE HOJA: CxP (CUENTAS POR PAGAR)
# ==============================================================================

def crear_hoja_cxp(wb):
    """Crea la hoja de CxP con fórmulas que buscan PENDIENTES"""
    print("📝 Creando hoja CxP (Cuentas por Pagar)...")

    # Crear hoja
    if 'CxP' in wb.sheetnames:
        ws = wb['CxP']
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('CxP')

    # Título
    ws['A1'] = '💰 CUENTAS POR PAGAR'
    ws['A1'].font = Font(bold=True, size=14, color='FF0000')

    # Resumen
    ws['A3'] = 'Total Pendiente:'
    ws['B3'] = '=SUM(D6:D55)'
    ws['B3'].number_format = '₡#,##0.00'
    ws['B3'].font = Font(bold=True, size=12)

    # Headers (fila 5)
    for nombre_campo, config in ESTRUCTURA_CXP.items():
        col = config['col']
        ws.cell(5, col, nombre_campo)

    # Aplicar estilo a headers
    aplicar_estilo_header(ws, 5, HOJAS_CONFIG['CxP']['color_header'])

    # Ajustar anchos
    ajustar_anchos_columnas(ws, ESTRUCTURA_CXP)

    # Congelar headers
    ws.freeze_panes = 'A6'

    # Insertar fórmulas (filas 6 a 55 = 50 filas)
    max_fila_trans = 1000  # Asumir máximo 1000 transacciones

    for fila in range(6, 56):
        formulas = get_formula_cxp(fila, max_fila_trans)

        for nombre_campo, config in ESTRUCTURA_CXP.items():
            col = config['col']
            formula = formulas.get(nombre_campo)

            if formula:
                ws.cell(fila, col, formula)

                # Aplicar formato si es necesario
                formato = config.get('formato')
                if formato:
                    ws.cell(fila, col).number_format = formato

    print("  ✅ Hoja CxP creada con fórmulas")
    return ws

# ==============================================================================
# CREACIÓN DE HOJA: CxC (CUENTAS POR COBRAR)
# ==============================================================================

def crear_hoja_cxc(wb):
    """Crea la hoja de CxC con fórmulas que buscan POR COBRAR"""
    print("📝 Creando hoja CxC (Cuentas por Cobrar)...")

    # Crear hoja
    if 'CxC' in wb.sheetnames:
        ws = wb['CxC']
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('CxC')

    # Título
    ws['A1'] = '💵 CUENTAS POR COBRAR'
    ws['A1'].font = Font(bold=True, size=14, color='00B050')

    # Resumen
    ws['A3'] = 'Total por Cobrar:'
    ws['B3'] = '=SUM(D6:D55)'
    ws['B3'].number_format = '₡#,##0.00'
    ws['B3'].font = Font(bold=True, size=12)

    # Headers (fila 5)
    for nombre_campo, config in ESTRUCTURA_CXC.items():
        col = config['col']
        ws.cell(5, col, nombre_campo)

    # Aplicar estilo
    aplicar_estilo_header(ws, 5, HOJAS_CONFIG['CxC']['color_header'])

    # Ajustar anchos
    ajustar_anchos_columnas(ws, ESTRUCTURA_CXC)

    # Congelar headers
    ws.freeze_panes = 'A6'

    # Insertar fórmulas
    max_fila_trans = 1000

    for fila in range(6, 56):
        formulas = get_formula_cxc(fila, max_fila_trans)

        for nombre_campo, config in ESTRUCTURA_CXC.items():
            col = config['col']
            formula = formulas.get(nombre_campo)

            if formula:
                ws.cell(fila, col, formula)

                formato = config.get('formato')
                if formato:
                    ws.cell(fila, col).number_format = formato

    print("  ✅ Hoja CxC creada con fórmulas")
    return ws

# ==============================================================================
# CREACIÓN DE HOJA: RESUMEN
# ==============================================================================

def crear_hoja_resumen(wb):
    """Crea la hoja de RESUMEN con dashboard"""
    print("📝 Creando hoja RESUMEN...")

    # Crear hoja
    if 'RESUMEN' in wb.sheetnames:
        ws = wb['RESUMEN']
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('RESUMEN', 0)  # Primera posición

    # Título
    ws['A1'] = '📊 DASHBOARD FINANCIERO v4.0'
    ws['A1'].font = Font(bold=True, size=16, color='7030A0')
    ws.merge_cells('A1:D1')

    # Métricas principales
    ws['A3'] = 'MÉTRICAS PRINCIPALES'
    ws['A3'].font = Font(bold=True, size=12)

    metricas = [
        ('A5', 'Total Ingresos:', 'B5', '=SUMIF(TRANSACCIONES!B:B,"INGRESO",TRANSACCIONES!I:I)'),
        ('A6', 'Total Egresos:', 'B6', '=SUMIF(TRANSACCIONES!B:B,"EGRESO",TRANSACCIONES!I:I)'),
        ('A7', 'Balance:', 'B7', '=B5-B6'),
        ('A9', 'Total por Pagar:', 'B9', '=CxP!B3'),
        ('A10', 'Total por Cobrar:', 'B10', '=CxC!B3'),
        ('A12', 'Transacciones Totales:', 'B12', '=COUNTA(TRANSACCIONES!A:A)-1'),
    ]

    for label_cell, label, value_cell, formula in metricas:
        ws[label_cell] = label
        ws[label_cell].font = Font(bold=True)
        ws[value_cell] = formula
        ws[value_cell].number_format = '₡#,##0.00' if 'SUM' in formula else '0'
        ws[value_cell].font = Font(bold=True, size=11)

    # Estado del sistema
    ws['A14'] = 'ESTADO DEL SISTEMA'
    ws['A14'].font = Font(bold=True, size=12)

    ws['A16'] = 'Versión:'
    ws['B16'] = VERSION
    ws['A17'] = 'Última actualización:'
    ws['B17'] = '=NOW()'
    ws['B17'].number_format = 'DD/MM/YYYY HH:MM'

    # Ajustar anchos
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20

    print("  ✅ Hoja RESUMEN creada")
    return ws

# ==============================================================================
# GENERACIÓN PRINCIPAL
# ==============================================================================

def generar_excel_v4(nombre_archivo='AlvaroVelasco_Finanzas_v4.0.xlsx', incluir_ejemplos=True):
    """
    Genera el archivo Excel v4.0 completo desde cero

    Args:
        nombre_archivo: Nombre del archivo a crear
        incluir_ejemplos: Si True, incluye transacciones de ejemplo

    Returns:
        str: Path del archivo creado
    """
    print(f"""
╔═══════════════════════════════════════════════════════════════════╗
║           GENERADOR DE EXCEL v4.0 - INICIO                        ║
╠═══════════════════════════════════════════════════════════════════╣
║ Archivo: {nombre_archivo:<55} ║
║ Incluir ejemplos: {'SI' if incluir_ejemplos else 'NO':<49} ║
╚═══════════════════════════════════════════════════════════════════╝
    """)

    # Crear workbook
    wb = openpyxl.Workbook()

    # Eliminar hoja por defecto
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    # Crear hojas en orden
    crear_hoja_transacciones(wb, incluir_ejemplos)
    crear_hoja_cxp(wb)
    crear_hoja_cxc(wb)
    crear_hoja_resumen(wb)

    # Activar hoja RESUMEN por defecto
    wb.active = wb['RESUMEN']

    # Guardar
    print(f"\n💾 Guardando archivo: {nombre_archivo}")
    wb.save(nombre_archivo)
    wb.close()

    print(f"""
╔═══════════════════════════════════════════════════════════════════╗
║           ✅ ARCHIVO EXCEL v4.0 GENERADO EXITOSAMENTE              ║
╠═══════════════════════════════════════════════════════════════════╣
║ Archivo creado: {nombre_archivo:<49} ║
║                                                                   ║
║ Hojas creadas:                                                    ║
║  ✅ RESUMEN (Dashboard)                                            ║
║  ✅ TRANSACCIONES (Registro completo)                              ║
║  ✅ CxP (Cuentas por Pagar con fórmulas)                           ║
║  ✅ CxC (Cuentas por Cobrar con fórmulas)                          ║
║                                                                   ║
║ Próximos pasos:                                                   ║
║  1. Abrir archivo en Excel                                        ║
║  2. Verificar que CxP muestra los pendientes                      ║
║  3. Verificar que CxC muestra los por cobrar                      ║
║  4. Ejecutar auditoría: python auditoria.py {nombre_archivo:<23} ║
╚═══════════════════════════════════════════════════════════════════╝
    """)

    return nombre_archivo

# ==============================================================================
# EJECUCIÓN PRINCIPAL
# ==============================================================================

if __name__ == '__main__':
    import sys
    import os

    # Cambiar al directorio raíz del proyecto
    os.chdir('..')

    # Parámetros
    nombre_archivo = 'AlvaroVelasco_Finanzas_v4.0.xlsx'
    incluir_ejemplos = True

    # Argumentos de línea de comandos
    if len(sys.argv) > 1:
        nombre_archivo = sys.argv[1]

    if len(sys.argv) > 2:
        incluir_ejemplos = sys.argv[2].lower() in ['true', '1', 'si', 'yes']

    # Generar
    try:
        archivo_creado = generar_excel_v4(nombre_archivo, incluir_ejemplos)

        print("\n🎉 ¡LISTO! El archivo está listo para usar.")
        print(f"📁 Ubicación: {os.path.abspath(archivo_creado)}")

        sys.exit(0)

    except Exception as e:
        print(f"\n❌ ERROR al generar archivo: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
