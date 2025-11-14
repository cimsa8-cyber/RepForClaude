# -*- coding: utf-8 -*-
"""
🔍 MÓDULO DE VALIDACIONES - Sistema de Finanzas v4.0
=====================================================

Funciones para validar la integridad del archivo Excel.

LECCIÓN v3.0:
- ❌ No validar estructura antes de escribir
- ✅ Validar headers, formatos, fórmulas y datos
"""

import openpyxl
from datetime import datetime
from config import (
    ESTRUCTURA_TRANSACCIONES,
    ESTRUCTURA_CXP,
    ESTRUCTURA_CXC,
    HOJAS_CONFIG,
    get_columna_por_nombre
)

# ==============================================================================
# VALIDACIÓN DE ESTRUCTURA
# ==============================================================================

def validar_estructura_hoja(ws, nombre_hoja):
    """
    Valida que los headers de una hoja coincidan con la estructura definida

    Args:
        ws: Worksheet de openpyxl
        nombre_hoja: Nombre de la hoja ('TRANSACCIONES', 'CxP', 'CxC')

    Returns:
        (bool, list): (es_valida, lista_errores)

    CRÍTICO: Evita el error v3.0 de columnas desalineadas
    """
    errores = []

    if nombre_hoja not in HOJAS_CONFIG:
        errores.append(f"Hoja '{nombre_hoja}' no está definida en configuración")
        return False, errores

    estructura = HOJAS_CONFIG[nombre_hoja]['estructura']

    if not estructura:  # RESUMEN no tiene estructura definida
        return True, []

    # Validar cada campo
    for nombre_campo, config in estructura.items():
        col = config['col']
        header_esperado = nombre_campo
        header_real = ws.cell(1, col).value

        if header_real != header_esperado:
            errores.append(
                f"❌ Columna {col} ({config.get('col_letra', '?')}): "
                f"Esperaba '{header_esperado}', encontró '{header_real}'"
            )

    if errores:
        return False, errores
    else:
        return True, []

def validar_todas_las_estructuras(wb):
    """
    Valida la estructura de todas las hojas del workbook

    Returns:
        (bool, dict): (es_valido, dict_errores_por_hoja)
    """
    errores_por_hoja = {}
    es_valido = True

    for nombre_hoja in ['TRANSACCIONES', 'CxP', 'CxC']:
        if nombre_hoja not in wb.sheetnames:
            errores_por_hoja[nombre_hoja] = [f"❌ Hoja '{nombre_hoja}' no existe"]
            es_valido = False
            continue

        ws = wb[nombre_hoja]
        valida, errores = validar_estructura_hoja(ws, nombre_hoja)

        if not valida:
            errores_por_hoja[nombre_hoja] = errores
            es_valido = False

    return es_valido, errores_por_hoja

# ==============================================================================
# VALIDACIÓN DE FORMATOS
# ==============================================================================

def validar_formatos_fecha(ws, nombre_hoja='TRANSACCIONES'):
    """
    Valida que las columnas de fecha tengan formato DD/MM/YYYY

    LECCIÓN v3.0: Las fechas se guardaban en formato ISO sin formato de celda
    """
    errores = []
    estructura = HOJAS_CONFIG[nombre_hoja]['estructura']

    # Buscar campos tipo datetime
    campos_fecha = {
        nombre: config for nombre, config in estructura.items()
        if config.get('tipo') == 'datetime'
    }

    if not campos_fecha:
        return True, []

    # Validar formato en primeras 10 filas con datos
    fila_inicio = HOJAS_CONFIG[nombre_hoja]['fila_inicio_datos']
    max_fila = min(fila_inicio + 10, ws.max_row + 1)

    for nombre_campo, config in campos_fecha.items():
        col = config['col']
        formato_esperado = config.get('formato_openpyxl', 'DD/MM/YYYY')

        for fila in range(fila_inicio, max_fila):
            celda = ws.cell(fila, col)
            if celda.value and isinstance(celda.value, datetime):
                if celda.number_format not in ['DD/MM/YYYY', 'dd/mm/yyyy']:
                    errores.append(
                        f"⚠️ {nombre_campo} fila {fila}: formato '{celda.number_format}' "
                        f"(esperado: {formato_esperado})"
                    )
                    break  # Solo reportar una vez por columna

    if errores:
        return False, errores
    else:
        return True, []

def validar_formatos_moneda(ws, nombre_hoja='TRANSACCIONES'):
    """Valida que las columnas de monto tengan formato de moneda"""
    errores = []
    estructura = HOJAS_CONFIG[nombre_hoja]['estructura']

    # Buscar campos tipo float con formato moneda
    campos_moneda = {
        nombre: config for nombre, config in estructura.items()
        if config.get('tipo') == 'float' and 'formato_openpyxl' in config
    }

    if not campos_moneda:
        return True, []

    fila_inicio = HOJAS_CONFIG[nombre_hoja]['fila_inicio_datos']
    max_fila = min(fila_inicio + 10, ws.max_row + 1)

    for nombre_campo, config in campos_moneda.items():
        col = config['col']
        formato_esperado = config.get('formato_openpyxl')

        for fila in range(fila_inicio, max_fila):
            celda = ws.cell(fila, col)
            if celda.value and isinstance(celda.value, (int, float)):
                if formato_esperado not in celda.number_format:
                    errores.append(
                        f"⚠️ {nombre_campo} fila {fila}: formato '{celda.number_format}' "
                        f"(esperado: {formato_esperado})"
                    )
                    break

    if errores:
        return False, errores
    else:
        return True, []

# ==============================================================================
# VALIDACIÓN DE DATOS
# ==============================================================================

def validar_transacciones_tienen_datos_minimos(ws):
    """
    Valida que cada transacción tenga al menos: Fecha, Tipo, Descripción, Monto, Estado

    Returns:
        (bool, list): (es_valida, lista_errores)
    """
    errores = []
    campos_requeridos = ['Fecha', 'Tipo', 'Descripción', 'Monto', 'Estado']

    fila_inicio = HOJAS_CONFIG['TRANSACCIONES']['fila_inicio_datos']

    for fila in range(fila_inicio, ws.max_row + 1):
        # Verificar si la fila tiene algún dato
        tiene_datos = any(ws.cell(fila, col).value for col in range(1, 16))

        if not tiene_datos:
            continue  # Fila vacía, ok

        # Si tiene datos, validar campos requeridos
        for campo in campos_requeridos:
            col = get_columna_por_nombre(campo, 'TRANSACCIONES')
            valor = ws.cell(fila, col).value

            if not valor:
                errores.append(
                    f"❌ Fila {fila}: falta campo requerido '{campo}'"
                )

    if errores:
        return False, errores
    else:
        return True, []

def validar_estados_validos(ws):
    """Valida que el campo Estado tenga valores válidos"""
    errores = []
    estados_validos = ['PAGADO', 'PENDIENTE', 'POR COBRAR', 'CANCELADO', 'PARCIAL']

    col_estado = get_columna_por_nombre('Estado', 'TRANSACCIONES')
    fila_inicio = HOJAS_CONFIG['TRANSACCIONES']['fila_inicio_datos']

    for fila in range(fila_inicio, ws.max_row + 1):
        estado = ws.cell(fila, col_estado).value

        if estado and str(estado).strip().upper() not in estados_validos:
            errores.append(
                f"❌ Fila {fila}: Estado '{estado}' no válido. "
                f"Válidos: {', '.join(estados_validos)}"
            )

    if errores:
        return False, errores
    else:
        return True, []

def validar_no_duplicados(ws):
    """
    Valida que no haya transacciones duplicadas

    Criterio: Misma fecha + descripción + monto = duplicado
    """
    errores = []
    transacciones_vistas = set()

    col_fecha = get_columna_por_nombre('Fecha', 'TRANSACCIONES')
    col_desc = get_columna_por_nombre('Descripción', 'TRANSACCIONES')
    col_monto = get_columna_por_nombre('Monto', 'TRANSACCIONES')

    fila_inicio = HOJAS_CONFIG['TRANSACCIONES']['fila_inicio_datos']

    for fila in range(fila_inicio, ws.max_row + 1):
        fecha = ws.cell(fila, col_fecha).value
        desc = ws.cell(fila, col_desc).value
        monto = ws.cell(fila, col_monto).value

        if not fecha or not desc:
            continue

        # Crear clave única
        clave = (str(fecha), str(desc), str(monto))

        if clave in transacciones_vistas:
            errores.append(
                f"⚠️ Fila {fila}: Posible duplicado (Fecha: {fecha}, "
                f"Descripción: {desc}, Monto: {monto})"
            )
        else:
            transacciones_vistas.add(clave)

    if errores:
        return False, errores
    else:
        return True, []

# ==============================================================================
# VALIDACIÓN DE CxP Y CxC
# ==============================================================================

def validar_cxp_tiene_datos(wb):
    """
    Valida que CxP muestre datos si existen transacciones PENDIENTES

    CRÍTICO: Este fue el error principal en v3.0
    """
    ws_trans = wb['TRANSACCIONES']
    ws_cxp = wb['CxP']

    # Contar transacciones PENDIENTES
    col_estado = get_columna_por_nombre('Estado', 'TRANSACCIONES')
    fila_inicio_trans = HOJAS_CONFIG['TRANSACCIONES']['fila_inicio_datos']

    pendientes = []
    for fila in range(fila_inicio_trans, ws_trans.max_row + 1):
        estado = ws_trans.cell(fila, col_estado).value
        if estado and 'PENDIENTE' in str(estado).upper():
            desc = ws_trans.cell(fila, get_columna_por_nombre('Descripción', 'TRANSACCIONES')).value
            monto = ws_trans.cell(fila, get_columna_por_nombre('Monto', 'TRANSACCIONES')).value
            pendientes.append((fila, desc, monto))

    # Contar filas con datos en CxP
    fila_inicio_cxp = HOJAS_CONFIG['CxP']['fila_inicio_datos']
    filas_cxp = 0

    for fila in range(fila_inicio_cxp, fila_inicio_cxp + 50):  # Revisar primeras 50 filas
        fecha = ws_cxp.cell(fila, 1).value
        if fecha and str(fecha).strip():
            filas_cxp += 1

    # Validar
    if len(pendientes) > 0 and filas_cxp == 0:
        return False, [
            f"❌ ERROR CRÍTICO: Hay {len(pendientes)} transacciones PENDIENTES "
            f"en TRANSACCIONES pero CxP está vacía!",
            "Transacciones pendientes encontradas:",
            *[f"  - Fila {f}: {d} (₡{m:,.2f})" for f, d, m in pendientes[:5]]
        ]
    elif len(pendientes) > filas_cxp:
        return False, [
            f"⚠️ ADVERTENCIA: Hay {len(pendientes)} PENDIENTES pero CxP solo muestra {filas_cxp} filas"
        ]
    else:
        return True, []

def validar_cxc_tiene_datos(wb):
    """Valida que CxC muestre datos si existen transacciones POR COBRAR"""
    ws_trans = wb['TRANSACCIONES']
    ws_cxc = wb['CxC']

    # Contar transacciones POR COBRAR
    col_estado = get_columna_por_nombre('Estado', 'TRANSACCIONES')
    fila_inicio_trans = HOJAS_CONFIG['TRANSACCIONES']['fila_inicio_datos']

    por_cobrar = []
    for fila in range(fila_inicio_trans, ws_trans.max_row + 1):
        estado = ws_trans.cell(fila, col_estado).value
        if estado and 'POR COBRAR' in str(estado).upper():
            desc = ws_trans.cell(fila, get_columna_por_nombre('Descripción', 'TRANSACCIONES')).value
            monto = ws_trans.cell(fila, get_columna_por_nombre('Monto', 'TRANSACCIONES')).value
            por_cobrar.append((fila, desc, monto))

    # Contar filas con datos en CxC
    fila_inicio_cxc = HOJAS_CONFIG['CxC']['fila_inicio_datos']
    filas_cxc = 0

    for fila in range(fila_inicio_cxc, fila_inicio_cxc + 50):
        fecha = ws_cxc.cell(fila, 1).value
        if fecha and str(fecha).strip():
            filas_cxc += 1

    # Validar
    if len(por_cobrar) > 0 and filas_cxc == 0:
        return False, [
            f"❌ ERROR CRÍTICO: Hay {len(por_cobrar)} transacciones POR COBRAR "
            f"en TRANSACCIONES pero CxC está vacía!",
            "Transacciones por cobrar encontradas:",
            *[f"  - Fila {f}: {d} (₡{m:,.2f})" for f, d, m in por_cobrar[:5]]
        ]
    elif len(por_cobrar) > filas_cxc:
        return False, [
            f"⚠️ ADVERTENCIA: Hay {len(por_cobrar)} POR COBRAR pero CxC solo muestra {filas_cxc} filas"
        ]
    else:
        return True, []

# ==============================================================================
# AUDITORÍA COMPLETA
# ==============================================================================

def auditoria_completa(archivo_excel):
    """
    Ejecuta todas las validaciones y retorna un reporte completo

    Returns:
        (bool, dict): (es_valido, reporte_completo)
    """
    print(f"\n🔍 Iniciando auditoría de: {archivo_excel}")
    print("=" * 70)

    try:
        wb = openpyxl.load_workbook(archivo_excel, data_only=False)
    except Exception as e:
        return False, {'error': f"No se pudo abrir el archivo: {e}"}

    reporte = {
        'archivo': archivo_excel,
        'fecha_auditoria': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'validaciones': {},
        'errores_criticos': 0,
        'advertencias': 0,
        'aprobado': True
    }

    # Lista de validaciones a ejecutar
    validaciones = [
        ('Estructura TRANSACCIONES', lambda: validar_estructura_hoja(wb['TRANSACCIONES'], 'TRANSACCIONES')),
        ('Estructura CxP', lambda: validar_estructura_hoja(wb['CxP'], 'CxP')),
        ('Estructura CxC', lambda: validar_estructura_hoja(wb['CxC'], 'CxC')),
        ('Formatos de fecha', lambda: validar_formatos_fecha(wb['TRANSACCIONES'])),
        ('Formatos de moneda', lambda: validar_formatos_moneda(wb['TRANSACCIONES'])),
        ('Datos mínimos en transacciones', lambda: validar_transacciones_tienen_datos_minimos(wb['TRANSACCIONES'])),
        ('Estados válidos', lambda: validar_estados_validos(wb['TRANSACCIONES'])),
        ('No duplicados', lambda: validar_no_duplicados(wb['TRANSACCIONES'])),
        ('CxP funcional', lambda: validar_cxp_tiene_datos(wb)),
        ('CxC funcional', lambda: validar_cxc_tiene_datos(wb))
    ]

    # Ejecutar validaciones
    for nombre, validacion in validaciones:
        try:
            es_valida, errores = validacion()

            reporte['validaciones'][nombre] = {
                'resultado': 'APROBADO' if es_valida else 'FALLIDO',
                'errores': errores
            }

            if not es_valida:
                if 'CRÍTICO' in str(errores):
                    reporte['errores_criticos'] += 1
                    reporte['aprobado'] = False
                else:
                    reporte['advertencias'] += 1

        except Exception as e:
            reporte['validaciones'][nombre] = {
                'resultado': 'ERROR',
                'errores': [f"Excepción: {e}"]
            }
            reporte['errores_criticos'] += 1
            reporte['aprobado'] = False

    wb.close()
    return reporte['aprobado'], reporte

def imprimir_reporte_auditoria(reporte):
    """Imprime el reporte de auditoría de forma legible"""
    print("\n╔═══════════════════════════════════════════════════════════════════╗")
    print("║                   REPORTE DE AUDITORÍA v4.0                       ║")
    print("╠═══════════════════════════════════════════════════════════════════╣")
    print(f"║ Archivo: {reporte['archivo']:<55} ║")
    print(f"║ Fecha: {reporte['fecha_auditoria']:<57} ║")
    print("╠═══════════════════════════════════════════════════════════════════╣")

    for nombre, resultado in reporte['validaciones'].items():
        estado = resultado['resultado']
        icono = '✅' if estado == 'APROBADO' else '❌'
        print(f"║ {icono} {nombre:<60} ║")

        if resultado['errores']:
            for error in resultado['errores'][:3]:  # Mostrar máximo 3 errores por validación
                print(f"║    {error:<63} ║")

    print("╠═══════════════════════════════════════════════════════════════════╣")
    print(f"║ Errores críticos: {reporte['errores_criticos']:<49} ║")
    print(f"║ Advertencias: {reporte['advertencias']:<53} ║")

    if reporte['aprobado']:
        print("║ Estado: ✅ APROBADO                                               ║")
    else:
        print("║ Estado: ❌ REQUIERE CORRECCIÓN                                    ║")

    print("╚═══════════════════════════════════════════════════════════════════╝\n")

# ==============================================================================
# EJECUCIÓN DIRECTA
# ==============================================================================

if __name__ == '__main__':
    import sys

    if len(sys.argv) < 2:
        print("Uso: python validaciones.py <archivo_excel>")
        sys.exit(1)

    archivo = sys.argv[1]
    aprobado, reporte = auditoria_completa(archivo)
    imprimir_reporte_auditoria(reporte)

    sys.exit(0 if aprobado else 1)
